# Source: https://www.newyorkfed.org/research/policy/rstar
"""Refresh ONLY US real-time HLW r-star from the official NY Fed vintage workbook.

Fails closed on changed layout, missing values, mismatched quarter, or a download failure.
Does not substitute current-model, LW, or third-party estimates.
"""
import argparse
import calendar
import datetime as dt
import hashlib
import io
import json
import math
import os
import re
import tempfile
import time
import urllib.request
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

OFFICIAL_URL = (
    "https://www.newyorkfed.org/medialibrary/media/research/economists/"
    "williams/data/Holston_Laubach_Williams_real_time_estimates.xlsx"
)
SHEET_NAME = re.compile(r"^(20\d{2})Q([1-4])$")
TIMEZONE = ZoneInfo("Asia/Taipei")


def download_official(url=OFFICIAL_URL):
    for attempt in range(3):
        try:
            request = urllib.request.Request(
                url,
                headers={
                    "User-Agent": "Mozilla/5.0 (compatible; LivingWaterAI HLW data verification)",
                    "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*",
                },
            )
            with urllib.request.urlopen(request, timeout=35) as response:
                payload = response.read(12000000)
            if not payload.startswith(b"PK\x03\x04"):
                raise ValueError("NY Fed did not return an xlsx ZIP workbook")
            return payload
        except Exception:
            if attempt == 2:
                raise
            time.sleep(2 * (attempt + 1))
    raise AssertionError("unreachable")


def get_real_time_us_rstar(payload):
    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    dated_sheets = sorted(
        (name for name in workbook.sheetnames if SHEET_NAME.fullmatch(name)),
        key=lambda name: (int(name[:4]), int(name[-1])),
    )
    if not dated_sheets:
        raise ValueError("No recognized quarterly real-time vintages in official workbook")
    sheet_name = dated_sheets[-1]
    ws = workbook[sheet_name]
    if ws["K5"].value != "Natural Rate (r*)" or ws["K6"].value != "US":
        raise ValueError("HLW real-time US r-star K-column headers changed; manual review required")
    if str(ws["A3"].value).strip() != "Final data point: " + sheet_name:
        raise ValueError("Workbook vintage and final observation mismatch")
    year, quarter = int(sheet_name[:4]), int(sheet_name[-1])
    quarter_first_month = quarter * 3 - 2
    final_point = ws.cell(ws.max_row, 1).value
    if not isinstance(final_point, (dt.date, dt.datetime)) or (
        final_point.year, final_point.month
    ) != (year, quarter_first_month):
        raise ValueError("HLW final row A-date not the expected quarter start")
    value = ws.cell(ws.max_row, 11).value
    if not isinstance(value, (int, float)) or isinstance(value, bool) or (
        not math.isfinite(value) or not -10 < value < 10
    ):
        raise ValueError("Missing or implausible official US real-time r-star")
    obs_month = quarter * 3
    obs_date = dt.date(year, obs_month, calendar.monthrange(year, obs_month)[1])
    if obs_date > dt.date.today():
        raise ValueError("Future-dated HLW real-time data")
    return {
        "value": float(value),
        "sheet": sheet_name,
        "cell": "K" + str(ws.max_row),
        "observation_date": obs_date.isoformat(),
        "sha256": hashlib.sha256(payload).hexdigest(),
        "url": OFFICIAL_URL,
    }


def apply_verified_value(data, verified, now):
    if data.get("schema") != "V70.3/10h-eight-supplement-v3":
        raise ValueError("Unexpected supplement schema")
    row = data["items"]["m06_rstar"]
    if row.get("source_url") != OFFICIAL_URL:
        raise ValueError("Official real-time source changed; manual review required")
    if not isinstance(row.get("candidate_checks"), list) or len(row["candidate_checks"]) != 4:
        raise ValueError("Expected four original candidate-source checks")
    previous_observation = row.get("observation_date")
    row.update(
        value=verified["value"],
        unit="%",
        observation_date=verified["observation_date"],
        source_url=verified["url"],
        method=(
            "NY Fed official Holston-Laubach-Williams REAL-TIME vintage workbook; "
            "sheet " + verified["sheet"] + "; US Natural Rate (r*) "
            + verified["cell"] + "; official xlsx SHA256 " + verified["sha256"]
        ),
        status="OK",
        detail=(
            "Official NY Fed Real-time estimates, not retrospective Current model estimates. "
            "Latest vintage " + verified["sheet"] + ", final US r-star "
            + verified["cell"] + ", verified direct download, sha256 "
            + verified["sha256"] + ". Quarter-end normalized for YYYY-MM-DD contract."
        ),
        retrieval_time=now.isoformat(timespec="seconds"),
        vintage=verified["sheet"] + " real-time as-released, not current revised model",
    )
    if previous_observation != verified["observation_date"]:
        # A new vintage must not inherit the previous publication date.
        row["publication_date"] = None
    check = row["candidate_checks"][1]
    check.update(
        site="NY Fed HLW real-time estimates",
        url=verified["url"],
        metric_or_horizon="US HLW real-time r* " + verified["sheet"],
        status="FOUND",
        checked_at=now.isoformat(timespec="seconds"),
        value=verified["value"],
        unit="%",
        observation_date=verified["observation_date"],
        notes=(
            "Verified directly against the official workbook SHA256 "
            + verified["sha256"] + " sheet " + verified["sheet"]
            + ", " + verified["cell"] + "."
        ),
    )
    data["generated_at"] = now.isoformat(timespec="seconds")
    data["as_of_taipei"] = now.strftime("%Y-%m-%d %H:%M:%S")
    return data


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", default="data/8_macro_supplement_verified.json")
    parser.add_argument("--dry-run", action="store_true")
    options = parser.parse_args()
    verified = get_real_time_us_rstar(download_official())
    now = dt.datetime.now(TIMEZONE)
    source_path = Path(options.input)
    data = json.loads(source_path.read_text(encoding="utf-8"))
    apply_verified_value(data, verified, now)
    if not options.dry_run:
        source_path.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(
            "w", encoding="utf-8", dir=source_path.parent,
            prefix=".hlw-", suffix=".tmp", delete=False
        ) as temp:
            temp.write(json.dumps(data, ensure_ascii=False, indent=2) + "\n")
            name = temp.name
        try:
            os.replace(name, source_path)
        finally:
            if os.path.exists(name):
                os.unlink(name)
    print("OFFICIAL_HLW_REALTIME_PASS", json.dumps(verified, ensure_ascii=False), flush=True)
    print("SUPPLEMENT_UPDATED", not options.dry_run, source_path, flush=True)


if __name__ == "__main__":
    main()
