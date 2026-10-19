#!/usr/bin/env python3
# 最後修改日期時間（台灣時間）：2026/10/19 08:03｜檔案名稱：bank_forecast_research.py｜版本：五大機構季度研究圖表乾跑
"""Extract and chart the handoff workbook's bank forecast evidence as research.

The workbook's dates are evidence cutoffs, not necessarily exact publication
timestamps. Every emitted comparison row is labeled PIT_PENDING accordingly.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import tempfile
from datetime import date, datetime, timezone
from pathlib import Path

REGISTER_FIELDS = [
    "forecast_id", "institution", "forecast_target_index", "forecast_value",
    "forecast_target_date", "evidence_date", "publication_date", "forecast_horizon_days",
    "institution_department", "source_url", "source_note", "verification_status",
    "source_artifact_sha256", "source_sheet", "source_row", "retrieved_at_utc",
]
COMPARISON_FIELDS = [
    "observation_date", "actual_trading_date", "sp500_start_close",
    "sp500_forward_end_date", "sp500_forward_end_close", "sp500_close",
    "institution", "forecast_id", "forecast_evidence_date", "forecast_target_date",
    "forecast_value", "forecast_horizon_days", "forecast_implied_return_pct",
    "sp500_forward_3m_return_pct", "return_status", "verification_status", "source_url",
]
INSTITUTIONS = {"高盛", "摩根大通", "摩根士丹利", "美國銀行", "瑞銀"}


class ResearchDataError(ValueError):
    pass


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _iso(value) -> str:
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return value.isoformat()
    if value is None:
        return ""
    return date.fromisoformat(str(value)[:10]).isoformat()


def extract_register(workbook: Path, retrieved_at_utc: str) -> list[dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ResearchDataError("openpyxl is required for workbook extraction") from exc
    wb = load_workbook(workbook, data_only=True, read_only=True)
    if "五大機構來源" not in wb.sheetnames:
        raise ResearchDataError("Workbook is missing sheet '五大機構來源'")
    ws = wb["五大機構來源"]
    iterator = ws.iter_rows(values_only=True)
    header_row = None
    for row in iterator:
        if row and str(row[0]).strip() == "公開資料最遲可得日期":
            header_row = [str(v).strip() if v is not None else "" for v in row]
            break
    expected = ["公開資料最遲可得日期", "機構", "年底目標點數", "預測截止日", "機構部門", "資料來源連結", "附註"]
    if header_row != expected:
        raise ResearchDataError(f"Unexpected bank source schema: {header_row}")

    workbook_sha = sha256_file(workbook)
    records = []
    seen_ids = set()
    for row_number, row in enumerate(iterator, start=4):
        if not row or all(v is None for v in row):
            continue
        try:
            evidence_date = _iso(row[0])
            institution = str(row[1]).strip()
            value = float(row[2])
            target_date = _iso(row[3])
            department = str(row[4]).strip()
            source_url = str(row[5]).strip()
            note = str(row[6]).strip() if row[6] is not None else ""
        except (IndexError, ValueError, TypeError) as exc:
            raise ResearchDataError(f"Malformed bank forecast at source row {row_number}") from exc
        if institution not in INSTITUTIONS or not source_url or value <= 0:
            raise ResearchDataError(f"Invalid institution, target, or source at row {row_number}")
        if date.fromisoformat(target_date) <= date.fromisoformat(evidence_date):
            raise ResearchDataError(f"Target date must follow evidence date at row {row_number}")
        stable = "|".join((evidence_date, institution, f"{value:g}", target_date, department, source_url, note))
        forecast_id = hashlib.sha256(stable.encode("utf-8")).hexdigest()[:20]
        if forecast_id in seen_ids:
            continue
        seen_ids.add(forecast_id)
        records.append({
            "forecast_id": forecast_id, "institution": institution,
            "forecast_target_index": "S&P 500", "forecast_value": f"{value:g}",
            "forecast_target_date": target_date, "evidence_date": evidence_date,
            "publication_date": "", "forecast_horizon_days": str(
                (date.fromisoformat(target_date) - date.fromisoformat(evidence_date)).days),
            "institution_department": department, "source_url": source_url,
            "source_note": note, "verification_status": "RESEARCH_PIT_PENDING",
            "source_artifact_sha256": workbook_sha, "source_sheet": "五大機構來源",
            "source_row": str(row_number), "retrieved_at_utc": retrieved_at_utc,
        })
    if not records:
        raise ResearchDataError("No bank forecasts found in source sheet")
    return records


def extract_quarterly_observations(workbook: Path) -> list[dict[str, str]]:
    from openpyxl import load_workbook
    wb = load_workbook(workbook, data_only=True, read_only=True)
    if "季度回測資料" not in wb.sheetnames:
        raise ResearchDataError("Workbook is missing sheet '季度回測資料'")
    ws = wb["季度回測資料"]
    rows = ws.iter_rows(values_only=True)
    header = None
    for row in rows:
        if row and row[0] == "基準日":
            header = list(row)
            break
    if not header or header[:6] != ["基準日", "實際交易日", "S&P500收盤點", "3M後交易日", "3M後收盤點", "3M價格報酬 (%)"]:
        raise ResearchDataError("Unexpected quarterly backtest sheet schema")
    output = []
    for row in rows:
        if not row or row[0] is None:
            continue
        if row[0] == "資料說明":
            break
        try:
            observation_date = _iso(row[0])
        except ValueError as exc:
            raise ResearchDataError(f"Invalid quarterly observation date: {row[0]!r}") from exc
        output.append({
            "observation_date": observation_date, "actual_trading_date": _iso(row[1]),
            "sp500_close": float(row[2]), "sp500_forward_end_date": _iso(row[3]) if row[3] else "",
            "sp500_forward_end_close": "" if row[4] is None else f"{float(row[4]):.8f}".rstrip("0").rstrip("."),
            "sp500_forward_3m_return_pct":
                "" if row[5] is None else f"{float(row[5]):.8f}".rstrip("0").rstrip("."),
        })
        if row[5] is not None and (row[3] is None or row[4] is None):
            raise ResearchDataError(f"Realized 3M return is missing endpoint date/close at row {row}")
    if not output:
        raise ResearchDataError("No quarterly observations in source workbook")
    return output


def build_comparison(register, observations):
    by_institution = {name: [] for name in INSTITUTIONS}
    for forecast in register:
        by_institution[forecast["institution"]].append(forecast)
    output = []
    for obs in observations:
        observed_on = date.fromisoformat(obs["actual_trading_date"])
        for institution in sorted(INSTITUTIONS):
            available = [f for f in by_institution[institution]
                         if date.fromisoformat(f["evidence_date"]) <= observed_on
                         and date.fromisoformat(f["forecast_target_date"]) > observed_on]
            if not available:
                continue
            forecast = max(available, key=lambda f: (f["evidence_date"], f["forecast_target_date"]))
            horizon = (date.fromisoformat(forecast["forecast_target_date"]) - observed_on).days
            implied = (float(forecast["forecast_value"]) / float(obs["sp500_close"]) - 1) * 100
            actual = obs["sp500_forward_3m_return_pct"]
            output.append({
                "observation_date": obs["observation_date"], "actual_trading_date": obs["actual_trading_date"],
                "sp500_start_close": f"{float(obs['sp500_close']):.8f}".rstrip("0").rstrip("."),
                "sp500_forward_end_date": obs["sp500_forward_end_date"],
                "sp500_forward_end_close": obs["sp500_forward_end_close"],
                "sp500_close": f"{float(obs['sp500_close']):.8f}".rstrip("0").rstrip("."),
                "institution": institution, "forecast_id": forecast["forecast_id"],
                "forecast_evidence_date": forecast["evidence_date"],
                "forecast_target_date": forecast["forecast_target_date"],
                "forecast_value": forecast["forecast_value"], "forecast_horizon_days": str(horizon),
                "forecast_implied_return_pct": f"{implied:.8f}".rstrip("0").rstrip("."),
                "sp500_forward_3m_return_pct": actual,
                "return_status": "REALIZED" if actual else "PENDING",
                "verification_status": "RESEARCH_PIT_PENDING", "source_url": forecast["source_url"],
            })
    return output


def write_csv(path: Path, fields: list[str], rows: list[dict[str, str]]):
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent, text=True)
    try:
        with os.fdopen(fd, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(rows)
            f.flush()
            os.fsync(f.fileno())
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.unlink(tmp)


def read_existing_register(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        if not reader.fieldnames or not set(REGISTER_FIELDS).issubset(reader.fieldnames):
            raise ResearchDataError("Existing bank register has an invalid schema")
        rows = [{field: row[field] for field in REGISTER_FIELDS} for row in reader]
    ids = [r["forecast_id"] for r in rows]
    if len(ids) != len(set(ids)):
        raise ResearchDataError("Existing bank register has duplicate forecast_id values")
    return rows


def merge_register(existing: list[dict[str, str]], incoming: list[dict[str, str]]) -> list[dict[str, str]]:
    """Append new forecast versions while retaining original evidence lineage."""
    merged = {r["forecast_id"]: dict(r) for r in existing}
    for row in incoming:
        old = merged.get(row["forecast_id"])
        if old:
            identity_fields = [f for f in REGISTER_FIELDS if f not in
                               {"source_artifact_sha256", "source_row", "retrieved_at_utc"}]
            if any(old[f] != row[f] for f in identity_fields):
                raise ResearchDataError(f"Conflicting bank record for forecast_id {row['forecast_id']}")
            continue
        merged[row["forecast_id"]] = dict(row)
    return sorted(merged.values(), key=lambda r: (r["evidence_date"], r["institution"], r["forecast_id"]))


def plot_charts(comparison, observations, output_dir: Path, generated_at: str):
    try:
        import matplotlib.pyplot as plt
        from matplotlib import font_manager
    except ImportError as exc:
        raise ResearchDataError("matplotlib is required to render charts") from exc
    font_candidates = [Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
                       Path("/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc")]
    for font_path in font_candidates:
        if font_path.is_file():
            font_manager.fontManager.addfont(str(font_path))
            plt.rcParams["font.family"] = font_manager.FontProperties(fname=str(font_path)).get_name()
            break
    else:
        raise ResearchDataError("No Traditional Chinese CJK font found; install Noto Sans CJK or an equivalent font")
    plt.rcParams["axes.unicode_minus"] = False
    colors = {"高盛": "#d62728", "摩根大通": "#1f77b4", "摩根士丹利": "#9467bd",
              "美國銀行": "#ff7f0e", "瑞銀": "#2ca02c"}
    grouped = {name: [r for r in comparison if r["institution"] == name] for name in INSTITUTIONS}
    output_dir.mkdir(parents=True, exist_ok=True)
    disclaimer = "研究資料｜年底目標隱含報酬與觀察日後3個月實報期限不同；PIT待核，不代表準確率"
    for institution in sorted(INSTITUTIONS):
        rows = grouped[institution]
        fig, ax = plt.subplots(figsize=(12, 6.6))
        x = [date.fromisoformat(r["actual_trading_date"]) for r in rows]
        forecast = [float(r["forecast_implied_return_pct"]) for r in rows]
        actual_rows = [r for r in observations if r["sp500_forward_3m_return_pct"]]
        actual_x = [date.fromisoformat(r["actual_trading_date"]) for r in actual_rows]
        actual = [float(r["sp500_forward_3m_return_pct"]) for r in actual_rows]
        ax.scatter(x, forecast, s=54, color=colors[institution], label="年底目標隱含報酬（觀察日快照）", zorder=3)
        ax.plot(actual_x, actual, marker="s", linewidth=2.7, color="#222222", label="S&P 500 後三個月實際價格報酬")
        ax.axhline(0, color="#777777", linewidth=0.9)
        ax.set_title(f"{institution}｜年底目標與 S&P 500 後三個月實報（季度研究）", fontsize=15, weight="bold")
        ax.set_ylabel("報酬率（%）")
        ax.grid(True, alpha=0.25)
        ax.legend(loc="best", frameon=True)
        ax.text(0.01, -0.19, f"{disclaimer}\n來源：交接工作簿五大機構來源表；生成：{generated_at}",
                transform=ax.transAxes, fontsize=8, va="top")
        fig.autofmt_xdate()
        fig.tight_layout(rect=(0, 0.08, 1, 1))
        fig.savefig(output_dir / f"{institution}_年底目標_vs_SP500後3M_季度研究.png", dpi=180)
        plt.close(fig)

    fig, ax = plt.subplots(figsize=(14, 8))
    for institution in sorted(INSTITUTIONS):
        rows = grouped[institution]
        x = [date.fromisoformat(r["actual_trading_date"]) for r in rows]
        y = [float(r["forecast_implied_return_pct"]) for r in rows]
        ax.scatter(x, y, s=42, color=colors[institution], label=f"{institution} 年底目標隱含報酬")
    realized = [r for r in observations if r["sp500_forward_3m_return_pct"]]
    ax.plot([date.fromisoformat(r["actual_trading_date"]) for r in realized],
            [float(r["sp500_forward_3m_return_pct"]) for r in realized], marker="s", linewidth=3,
            color="#202020", label="S&P 500 後三個月實際價格報酬")
    ax.axhline(0, color="#777777", linewidth=0.9)
    ax.set_title("五大機構年底目標與 S&P 500 後三個月實報｜季度研究總覽", fontsize=16, weight="bold")
    ax.set_ylabel("報酬率（%）")
    ax.grid(True, alpha=0.25)
    ax.legend(loc="best", ncol=2, frameon=True)
    ax.text(0.01, -0.15, f"{disclaimer}\n28個季度錨點加2026-10-16截止觀察；發布日期及來源需逐筆複核。生成：{generated_at}",
            transform=ax.transAxes, fontsize=8, va="top")
    fig.autofmt_xdate()
    fig.tight_layout(rect=(0, 0.08, 1, 1))
    fig.savefig(output_dir / "五大機構與SP500後3M_季度研究總覽.png", dpi=180)
    plt.close(fig)


def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--workbook", required=True, type=Path)
    p.add_argument("--output-dir", required=True, type=Path)
    args = p.parse_args(argv)
    try:
        now = datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
        incoming = extract_register(args.workbook, now)
        register_path = args.output_dir / "bank_forecast_register_research.csv"
        register = merge_register(read_existing_register(register_path), incoming)
        observations = extract_quarterly_observations(args.workbook)
        comparison = build_comparison(register, observations)
        if not comparison:
            raise ResearchDataError("No point-in-time eligible bank comparisons")
        write_csv(register_path, REGISTER_FIELDS, register)
        write_csv(args.output_dir / "bank_quarterly_comparison_research.csv", COMPARISON_FIELDS, comparison)
        plot_charts(comparison, observations, args.output_dir / "charts", now)
    except (ResearchDataError, OSError, ValueError) as exc:
        p.error(str(exc))
    print(json.dumps({"forecast_records": len(register), "quarterly_observations": len(observations),
                      "comparison_rows": len(comparison), "charts": 6,
                      "output_dir": str(args.output_dir), "status": "RESEARCH_PIT_PENDING"}, ensure_ascii=False))


if __name__ == "__main__":
    main()
